"""
merge_trades_into_1m.py

Merges GHULAM-OFBOT's 110 trade entries into the 1-minute BTCUSDT
data file, matching each trade to its NEAREST minute candle (trade
timestamps have fractional-second precision; 1m candles don't).

- Only Entry Price is carried over from the trades file (Side is kept
  only to determine row color, not written as data).
- LONG entry rows are highlighted green, SHORT entry rows red.
- All 1m rows before the first trade's entry time are excluded.

USAGE:
    python merge_trades_into_1m.py
"""

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

IN_1M_PATH = r"D:\data\btc\BTCUSDT_1m_aug2026_latest.parquet"
OUT_PATH = r"D:\data\btc\BTCUSDT_1m_with_GHULAM_trades.xlsx"

FONT = "Arial"

TRADES = [
    ("2026-08-08 14:13:01.952", "LONG", 65033.55),
    ("2026-08-08 21:54:08.470", "LONG", 65027.55),
    ("2026-08-09 06:10:06.860", "LONG", 64789.91),
    ("2026-08-09 09:02:50.581", "SHORT", 64787.59),
    ("2026-08-09 15:42:02.130", "LONG", 65020.5),
    ("2026-08-10 04:14:09.527", "SHORT", 65002.35),
    ("2026-08-10 04:39:51.711", "LONG", 65051.26),
    ("2026-08-10 07:11:49.013", "LONG", 65270.2),
    ("2026-08-10 13:47:40.281", "LONG", 64659.18),
    ("2026-08-10 15:23:01.571", "LONG", 64500.45),
    ("2026-08-10 18:19:21.739", "SHORT", 64096.18),
    ("2026-08-10 20:35:49.185", "LONG", 64092.77),
    ("2026-08-10 22:44:27.283", "SHORT", 63994.2),
    ("2026-08-11 01:06:20.312", "LONG", 63992.25),
    ("2026-08-11 15:44:08.864", "LONG", 63571.86),
    ("2026-08-12 13:02:52.505", "LONG", 64102.77),
    ("2026-08-12 14:51:24.153", "SHORT", 63505.05),
    ("2026-08-13 11:37:31.157", "SHORT", 63343.98),
    ("2026-08-13 13:56:16.427", "LONG", 63753.4),
    ("2026-08-13 17:11:56.289", "SHORT", 63096.33),
    ("2026-08-14 11:55:12.463", "SHORT", 62883.67),
    ("2026-08-14 13:54:19.179", "LONG", 62652.78),
    ("2026-08-15 00:21:00.257", "SHORT", 62973.65),
    ("2026-08-15 00:29:39.400", "LONG", 63017.95),
    ("2026-08-15 00:47:20.568", "SHORT", 62985.95),
    ("2026-08-15 00:51:55.456", "LONG", 63013.15),
    ("2026-08-15 01:24:48.904", "SHORT", 62972.55),
    ("2026-08-15 05:47:14.446", "LONG", 63047.31),
    ("2026-08-16 17:51:12.125", "LONG", 63140.18),
    ("2026-08-17 17:36:04.318", "LONG", 64128.87),
    ("2026-08-18 14:40:24.847", "LONG", 64947.94),
    ("2026-08-19 01:00:48.309", "LONG", 64468.84),
    ("2026-08-19 13:59:59.082", "LONG", 65007.35),
    ("2026-08-19 14:09:48.419", "LONG", 65361.42),
    ("2026-08-19 16:38:02.630", "LONG", 68817.81),
    ("2026-08-19 17:35:28.953", "LONG", 68023.25),
    ("2026-08-19 18:18:16.507", "SHORT", 68164.71),
    ("2026-08-19 18:37:23.365", "LONG", 68287.1),
    ("2026-08-19 19:08:21.422", "SHORT", 68352.28),
    ("2026-08-19 19:34:25.149", "LONG", 68313.61),
    ("2026-08-19 20:03:45.066", "SHORT", 68373.87),
    ("2026-08-19 21:43:38.412", "LONG", 69595.87),
    ("2026-08-19 22:17:14.280", "LONG", 69334.71),
    ("2026-08-19 23:00:12.201", "SHORT", 69208.11),
    ("2026-08-19 23:34:32.208", "LONG", 69320.01),
    ("2026-08-20 01:00:25.190", "LONG", 69399.43),
    ("2026-08-20 01:44:12.048", "SHORT", 69555.64),
    ("2026-08-20 03:00:18.220", "LONG", 69309.31),
    ("2026-08-20 04:30:13.359", "SHORT", 69277.69),
    ("2026-08-20 04:39:51.019", "LONG", 69346.82),
    ("2026-08-20 08:13:08.579", "LONG", 71058.16),
    ("2026-08-20 08:35:26.169", "SHORT", 71150.22),
    ("2026-08-20 09:03:58.965", "LONG", 71697.49),
    ("2026-08-20 09:34:09.586", "SHORT", 71865.07),
    ("2026-08-20 09:37:15.426", "LONG", 71988.34),
    ("2026-08-20 10:39:02.677", "LONG", 71890.53),
    ("2026-08-20 10:51:43.819", "SHORT", 71810.68),
    ("2026-08-20 11:13:09.939", "LONG", 71924.43),
    ("2026-08-20 11:43:04.574", "SHORT", 72147.02),
    ("2026-08-20 12:00:53.833", "LONG", 71930.83),
    ("2026-08-20 12:38:05.805", "SHORT", 71714.0),
    ("2026-08-20 13:17:36.186", "SHORT", 71925.56),
    ("2026-08-20 20:12:48.474", "LONG", 72739.8),
    ("2026-08-21 00:12:20.417", "SHORT", 73227.0),
    ("2026-08-21 00:44:34.108", "LONG", 73749.0),
    ("2026-08-21 00:53:59.189", "LONG", 73750.0),
    ("2026-08-21 01:18:49.692", "LONG", 74614.97),
    ("2026-08-21 01:26:45.033", "SHORT", 74488.15),
    ("2026-08-21 01:35:22.797", "SHORT", 74703.01),
    ("2026-08-21 02:24:00.172", "LONG", 74799.01),
    ("2026-08-21 03:08:40.080", "LONG", 74347.92),
    ("2026-08-21 04:08:22.503", "SHORT", 74703.21),
    ("2026-08-21 05:37:30.209", "LONG", 75289.5),
    ("2026-08-21 05:47:30.807", "SHORT", 75125.12),
    ("2026-08-21 06:17:31.552", "LONG", 75420.33),
    ("2026-08-21 06:31:29.611", "SHORT", 75374.77),
    ("2026-08-21 07:03:16.993", "LONG", 75808.81),
    ("2026-08-21 07:49:31.459", "SHORT", 76332.68),
    ("2026-08-21 07:58:17.235", "LONG", 76392.63),
    ("2026-08-21 08:17:04.513", "LONG", 76995.35),
    ("2026-08-21 08:29:09.170", "LONG", 77105.37),
    ("2026-08-21 09:03:57.223", "LONG", 77923.53),
    ("2026-08-21 09:28:47.168", "SHORT", 77707.81),
    ("2026-08-21 10:50:51.494", "LONG", 77786.2),
    ("2026-08-21 11:28:57.121", "LONG", 77445.34),
    ("2026-08-21 12:23:20.804", "LONG", 76879.02),
    ("2026-08-21 12:46:33.548", "LONG", 77355.87),
    ("2026-08-21 14:55:04.328", "SHORT", 77231.4),
    ("2026-08-21 16:18:13.104", "LONG", 77186.18),
    ("2026-08-21 16:34:25.681", "SHORT", 77273.99),
    ("2026-08-21 18:02:28.399", "LONG", 77454.24),
    ("2026-08-21 19:11:53.779", "SHORT", 76989.25),
    ("2026-08-21 20:46:10.942", "LONG", 77447.34),
    ("2026-08-21 20:51:16.638", "SHORT", 77397.27),
    ("2026-08-21 21:44:13.535", "LONG", 78216.79),
    ("2026-08-21 22:09:05.542", "SHORT", 78558.64),
    ("2026-08-21 22:54:10.146", "SHORT", 78505.25),
    ("2026-08-21 23:14:13.184", "LONG", 78539.95),
    ("2026-08-21 23:47:54.687", "LONG", 78292.91),
    ("2026-08-22 00:23:44.727", "SHORT", 78009.44),
    ("2026-08-22 01:55:00.381", "SHORT", 77799.59),
    ("2026-08-22 02:47:20.795", "LONG", 78379.32),
    ("2026-08-22 03:27:20.562", "LONG", 78736.19),
    ("2026-08-22 05:21:25.376", "LONG", 77068.36),
    ("2026-08-22 05:37:49.835", "LONG", 77251.8),
    ("2026-08-22 06:16:01.158", "LONG", 77252.4),
    ("2026-08-22 08:17:27.317", "SHORT", 77290.19),
    ("2026-08-22 09:07:44.302", "LONG", 77142.98),
    ("2026-08-22 10:02:01.391", "LONG", 76999.35),
    ("2026-08-22 11:14:57.922", "LONG", 76979.64),
]


def main():
    print(f"Loading {IN_1M_PATH} ...")
    df = pd.read_parquet(IN_1M_PATH)
    df["timestamp"] = pd.to_datetime(df["timestamp"], utc=True).dt.tz_localize(None)
    print(f"{len(df):,} 1m candles loaded, {df['timestamp'].iloc[0]} to {df['timestamp'].iloc[-1]}")

    trades = pd.DataFrame(TRADES, columns=["entry_time", "side", "entry_price"])
    trades["entry_time"] = pd.to_datetime(trades["entry_time"])
    trades["rounded_minute"] = trades["entry_time"].dt.round("min")

    valid_minutes = set(df["timestamp"])
    missing = trades[~trades["rounded_minute"].isin(valid_minutes)]
    if len(missing) > 0:
        print(f"WARNING: {len(missing)} trade(s) rounded to a minute with no matching candle:")
        print(missing[["entry_time", "rounded_minute"]])

    trade_map = trades.set_index("rounded_minute")[["side", "entry_price"]]
    df = df.merge(trade_map, left_on="timestamp", right_index=True, how="left")

    first_trade_minute = trades["rounded_minute"].min()
    before_count = (df["timestamp"] < first_trade_minute).sum()
    df = df[df["timestamp"] >= first_trade_minute].reset_index(drop=True)
    print(f"Excluded {before_count:,} rows before the first trade's candle ({first_trade_minute})")
    print(f"{len(df):,} rows remain, containing {df['entry_price'].notna().sum()} matched trade entries "
          f"(expected {len(trades)})")

    df = df.rename(columns={
        "timestamp": "Time", "open": "Open", "high": "High", "low": "Low",
        "close": "Close", "volume": "Volume", "entry_price": "Trade Entry Price",
    })
    df = df.drop(columns=["taker_buy_base", "taker_buy_quote", "trades"], errors="ignore")

    print(f"Writing {OUT_PATH} ...")
    with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
        export_cols = [c for c in df.columns if c != "side"]
        df[export_cols].to_excel(writer, sheet_name="1m with GHULAM Trades", index=False)
        ws = writer.sheets["1m with GHULAM Trades"]

        header_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for c in range(1, len(export_cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        normal_font = Font(name=FONT, size=9)
        thin_border = Border(*[Side(style="thin", color="DDDDDD")] * 4)
        green_fill = PatternFill("solid", fgColor="D9EAD3")
        red_fill = PatternFill("solid", fgColor="F4CCCC")

        sides = df["side"].values
        for r in range(2, ws.max_row + 1):
            side_val = sides[r - 2]
            for c in range(1, len(export_cols) + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = normal_font
                cell.border = thin_border
                if side_val == "LONG":
                    cell.fill = green_fill
                elif side_val == "SHORT":
                    cell.fill = red_fill

        widths = {"Time": 20}
        for i, col in enumerate(export_cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)
        ws.freeze_panes = "A2"

    print(f"Done. {len(df):,} rows written to {OUT_PATH}")


if __name__ == "__main__":
    main()
