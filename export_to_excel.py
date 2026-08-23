"""
export_to_excel.py

Converts the 5m CMA-Signal parquet output (every score/component
column, matching the Pine script exactly) into a formatted Excel file.

USAGE:
    python export_to_excel.py
"""

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

IN_PATH = r"D:\data\btc\BTCUSDT_5m_CMA_signals.parquet"
OUT_PATH = r"D:\data\btc\BTCUSDT_5m_CMA_signals.xlsx"

FONT = "Arial"

def main():
    print(f"Loading {IN_PATH} ...")
    df = pd.read_parquet(IN_PATH)
    print(f"{len(df):,} rows, {len(df.columns)} columns")

    # Excel cannot store timezone-aware datetimes -- strip tz info (values
    # remain correct UTC wall-clock times, just without the tz marker Excel
    # can't represent)
    df["timestamp"] = pd.to_datetime(df["timestamp"]).dt.tz_localize(None)

    # Friendly column order/renaming for readability, keeping every
    # underlying value from the pipeline
    df = df.rename(columns={
        "timestamp": "Time",
        "open": "Open", "high": "High", "low": "Low", "close": "Close", "volume": "Volume",
        "trend_long_score": "Trend Long Score", "trend_short_score": "Trend Short Score",
        "momentum_long_score": "Momentum Long Score", "momentum_short_score": "Momentum Short Score",
        "volatility_score": "Volatility Score", "vol_regime": "Volatility Regime",
        "vwap_long_score": "VWAP Long Score", "vwap_short_score": "VWAP Short Score",
        "volume_score": "Volume Score",
        "structure_long_score": "Structure Long Score", "structure_short_score": "Structure Short Score",
        "long_score": "TOTAL LONG SCORE", "short_score": "TOTAL SHORT SCORE",
        "mtf_long_ok": "MTF Long OK", "mtf_short_ok": "MTF Short OK",
        "signal_state": "Signal State", "new_long_signal": "New LONG Signal", "new_short_signal": "New SHORT Signal",
    })

    print(f"Writing {OUT_PATH} ...")
    with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="5m CMA Signals", index=False)
        ws = writer.sheets["5m CMA Signals"]

        # Header styling
        header_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for c in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        # Body font + light borders
        normal_font = Font(name=FONT, size=9)
        thin_border = Border(*[Side(style="thin", color="DDDDDD")] * 4)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.font = normal_font
                cell.border = thin_border

        # Highlight actual signal rows for quick visual scanning
        green_fill = PatternFill("solid", fgColor="D9EAD3")
        red_fill = PatternFill("solid", fgColor="F4CCCC")
        long_col = df.columns.get_loc("New LONG Signal") + 1
        short_col = df.columns.get_loc("New SHORT Signal") + 1
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=long_col).value:
                for c in range(1, len(df.columns) + 1):
                    ws.cell(row=r, column=c).fill = green_fill
            elif ws.cell(row=r, column=short_col).value:
                for c in range(1, len(df.columns) + 1):
                    ws.cell(row=r, column=c).fill = red_fill

        # Reasonable column widths
        widths = {"Time": 20, "Volatility Regime": 14, "Signal State": 12}
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 11)

        ws.freeze_panes = "A2"

    print(f"Done. {len(df):,} rows written to {OUT_PATH}")
    n_long = df["New LONG Signal"].sum()
    n_short = df["New SHORT Signal"].sum()
    print(f"Signal rows highlighted: {n_long} LONG (green), {n_short} SHORT (red)")


if __name__ == "__main__":
    main()
