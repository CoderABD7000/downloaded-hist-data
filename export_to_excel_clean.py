"""
export_to_excel_clean.py

Same as export_to_excel.py, but WITHOUT the signal columns (Signal
State, New LONG Signal, New SHORT Signal) or the green/red row
highlighting -- just the clean OHLCV + every calculated parameter
column from the Pine script translation.

USAGE:
    python export_to_excel_clean.py
"""

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

IN_PATH = r"D:\data\btc\BTCUSDT_5m_CMA_signals.parquet"
OUT_PATH = r"D:\data\btc\BTCUSDT_5m_CMA_parameters_clean.xlsx"

FONT = "Arial"

def main():
    print(f"Loading {IN_PATH} ...")
    df = pd.read_parquet(IN_PATH)
    print(f"{len(df):,} rows, {len(df.columns)} columns")

    df["timestamp"] = pd.to_datetime(df["timestamp"]).dt.tz_localize(None)

    # Drop signal-specific columns -- keep only OHLCV + calculated parameters
    df = df.drop(columns=["signal_state", "new_long_signal", "new_short_signal"])

    df = df.rename(columns={
        "timestamp": "Time",
        "open": "Open", "high": "High", "low": "Low", "close": "Close", "volume": "Volume",
        "trend_long_score": "Trend Long Score", "trend_short_score": "Trend Short Score",
        "momentum_long_score": "Momentum Long Score", "momentum_short_score": "Momentum Short Score",
        "volatility_score": "Volatility Score", "vol_regime": "Volatility Regime",
        "vwap_long_score": "VWAP Long Score", "vwap_short_score": "VWAP Short Score",
        "volume_score": "Volume Score",
        "structure_long_score": "Structure Long Score", "structure_short_score": "Structure Short Score",
        "long_score": "TOTAL LONG SCORE", "short_score": "TOTAL SHORT SCORE",
        "mtf_long_ok": "MTF Long OK", "mtf_short_ok": "MTF Short OK",
    })

    print(f"Writing {OUT_PATH} ...")
    with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="5m CMA Parameters", index=False)
        ws = writer.sheets["5m CMA Parameters"]

        header_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for c in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        normal_font = Font(name=FONT, size=9)
        thin_border = Border(*[Side(style="thin", color="DDDDDD")] * 4)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.font = normal_font
                cell.border = thin_border

        widths = {"Time": 20, "Volatility Regime": 14}
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 11)

        ws.freeze_panes = "A2"

    print(f"Done. {len(df):,} rows, {len(df.columns)} columns written to {OUT_PATH}")


if __name__ == "__main__":
    main()
